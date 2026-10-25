"""Editorial PowerPoint composition. Content is supplied verbatim by the caller."""
from __future__ import annotations

import json
import math
import re
from pathlib import Path
from functools import lru_cache

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

BG = '0B0E14'
WHITE = 'F3F1EB'
MUTED = 'A5ADBA'
RULE = '303640'
AMBER = 'C9AD78'
W, H, M, CW, BOTTOM = 13.333333, 7.5, .65, 12.033333, 6.7


@lru_cache(maxsize=100)
def font(size, bold=False):
    path = Path('C:/Windows/Fonts') / ('seguisb.ttf' if bold else 'segoeui.ttf')
    return ImageFont.truetype(str(path), round(size * 4))


def text_height(text, width, size, bold=False, leading=1.18):
    """Conservative physical line estimate; Office rendering is the final check."""
    f = font(size, bold)
    available = max(1, (width * 72 - 9) * 4)
    count = 0
    for paragraph in text.replace('\v', '\n').split('\n'):
        line = ''
        for word in paragraph.split():
            candidate = f'{line} {word}'.strip()
            if line and f.getlength(candidate) > available:
                count += 1
                line = word
            else:
                line = candidate
        count += 1
    return count * size * leading / 72 + .055


def mixed_height(text, width, size, leading):
    available = (width * 72 - 9) * 4
    count = 0
    for paragraph in text.split('\n'):
        label_end = paragraph.find(':')
        used, offset = 0, 0
        for word in paragraph.split():
            weight = label_end >= 0 and offset <= label_end
            f = font(size, weight)
            amount = f.getlength(word + ' ')
            if used and used + amount > available:
                count += 1
                used = 0
            used += amount
            offset += len(word) + 1
        count += 1
    return count * size * leading / 72 + .055


def tx(slide, text, x, y, width, size=18.5, color=WHITE, bold=False,
       height=None, leading=1.18, label_bold=False, align=None):
    estimated = (mixed_height(text, width, size, leading) if label_bold else
                 text_height(text, width, size, bold, leading))
    height = estimated if height is None else height
    if x < 0 or y < 0 or x + width > W + .001 or y + height > H + .001:
        raise ValueError(f'Text outside canvas: {x, y, width, height}')
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, content in enumerate(text.replace('\v', '\n').split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = Pt(size * leading)
        if align is not None:
            p.alignment = align
        label_end = content.find(':') if label_bold else -1
        used = 0
        for token in re.finditer(r'\S+', content):
            value = token.group()
            weight = bold or (label_end >= 0 and token.start() <= label_end)
            ink = WHITE if label_bold and weight else color
            face = font(size, weight)
            separator = ' ' if used else ''
            amount = face.getlength(separator + value) / 4
            if used and used + amount > width * 72 - 9:
                p.add_line_break()
                used, separator = 0, ''
                amount = face.getlength(value) / 4
            r = p.add_run()
            r.text = separator + value
            r.font.name = 'Segoe UI Semibold' if weight else 'Segoe UI'
            r.font.size = Pt(size)
            r.font.bold = weight
            r.font.color.rgb = RGBColor.from_string(ink)
            used += amount
    return height


def rule(slide, y, x=M, width=CW, color=RULE, thickness=.7):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                   Inches(width), Pt(thickness))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()


class Composer:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.mapping = []

    def page(self, source, title, layout):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(BG)
        self.mapping.append({'slide_number': len(self.mapping) + 1,
                             'source_slide': source, 'title': title, 'layout': layout})
        return s

    def header(self, s, tag, title, subtitle=''):
        tx(s, tag, M, .43, CW, 10.5, AMBER, True)
        y = .76
        y += tx(s, title, M, y, CW, 30, bold=True, leading=1.10) + .12
        if subtitle:
            y += tx(s, subtitle, M, y, CW, 14.5, MUTED, leading=1.14) + .17
        else:
            y += .10
        rule(s, y)
        return y + .24

    def finish(self, source):
        for i, (s, record) in enumerate(zip(self.prs.slides, self.mapping)):
            rule(s, 6.99)
            tx(s, f"РТН  /  МАТЕРІАЛИ БРИФІНГУ", M, 7.08, 7.8, 9, MUTED)
            tx(s, f'{i+1:02} / {len(self.mapping):02}', 11.1, 7.07, 1.58,
               10, AMBER, align=PP_ALIGN.RIGHT)
            original = source.slides[record['source_slide'] - 1]
            strings = []
            for sh in original.shapes:
                if sh.has_text_frame and sh.text.strip():
                    strings.append(sh.text)
                elif sh.has_table:
                    strings.extend(c.text for row in sh.table.rows for c in row.cells)
            s.notes_slide.notes_text_frame.text = (
                f"Source slide: {record['source_slide']}\n"
                'Visual formatting of supplied material; content not independently verified.\n\n'
                + '\n\n'.join(strings))


def compose(data, source_path, output_path):
    output = Path(output_path).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    if output == Path(source_path).resolve():
        raise ValueError('Source and output must be different files')
    c = Composer()
    d = data['source_1']
    s = c.page(1, d['title'], 'cover')
    tx(s, d['tag'], M, .65, 10.5, 11, AMBER, True)
    rule(s, 1.35, width=1.1, color=AMBER, thickness=2)
    title_y = 1.78
    title_h = tx(s, d['title'], M, title_y, 10.65, 43, bold=True, leading=1.11)
    tx(s, d['subtitle'], M, title_y + title_h + .52, 9.2, 19, MUTED)
    tx(s, 'РТН', 10.9, 5.7, 1.78, 28, AMBER, True, align=PP_ALIGN.RIGHT)

    s = c.page(1, d['title'], 'overview_rows')
    y = c.header(s, d['tag'], 'Зміст брифінгу')
    for i, (title, desc) in enumerate(d['features']):
        tx(s, f'0{i+1}', M, y, .65, 21, AMBER)
        lh = tx(s, title, 1.65, y, 3.2, 18, bold=True)
        rh = tx(s, desc, 5.25, y, 7.43, 20, WHITE, leading=1.16)
        y += max(lh, rh) + .26
        if i < 2:
            rule(s, y)
            y += .25
    tx(s, d['footer'], M, 6.28, CW, 13.5, MUTED)

    d = data['source_2']
    for group_idx, group in enumerate((d['threats'][:3], d['threats'][3:])):
        s = c.page(2, d['title'], 'open_metric_rows')
        y = c.header(s, d['tag'], d['title'], d['subtitle'] if group_idx == 0 else '')
        for k, item in enumerate(group):
            name_h = text_height(item['title'], 6.7, 18.5, True)
            metric_h = text_height(item['metric'], 4.55, 24, True)
            desc_h = text_height(item['desc'], CW, 18)
            needed = max(name_h, metric_h) + .08 + desc_h
            if y + needed > BOTTOM:
                s = c.page(2, d['title'], 'open_metric_rows')
                y = c.header(s, d['tag'], d['title'])
            tx(s, item['title'], M, y, 6.7, 18.5, bold=True)
            tx(s, item['metric'], 8.13, y - .02, 4.55, 24, AMBER, True,
               align=PP_ALIGN.RIGHT)
            y += max(name_h, metric_h) + .08
            y += tx(s, item['desc'], M, y, CW, 18, MUTED) + .20
            if k < len(group) - 1:
                rule(s, y)
                y += .20

    # An oversized typographic metric repeats the original record without new claims.
    item = d['threats'][4]
    s = c.page(2, item['title'], 'hero_metric')
    c.header(s, d['tag'], item['title'], d['subtitle'])
    value, unit = item['metric'].split('/', 1)
    tx(s, value.strip(), M, 2.75, 6.7, 92, AMBER, leading=1.05)
    tx(s, '/ ' + unit.strip(), M + .08, 4.40, 5.8, 24, MUTED)
    tx(s, item['desc'], 8.0, 2.97, 4.68, 24, WHITE, leading=1.18)

    d = data['source_3']
    for start in (0, 2):
        s = c.page(3, d['title'], 'two_column_editorial')
        y = c.header(s, d['tag'], d['title'], d['subtitle'] if start == 0 else '')
        for col, block in enumerate(d['phases'][start:start+2]):
            x, width, yy = M + col * 6.30, 5.73, y
            yy += tx(s, f'{start+col+1:02}', x, yy, width, 46, AMBER) + .08
            yy += tx(s, block['title'], x, yy, width, 18, bold=True) + .10
            yy += tx(s, block['subtitle'], x, yy, width, 14.5, MUTED) + .17
            for item in block['bullets']:
                yy += tx(s, item, x, yy, width, 18, leading=1.12) + .10
            if yy > BOTTOM + .10:
                raise ValueError(f'Phase layout needs another page: {start+col+1} ({yy:.2f})')

    d = data['source_4']
    for start in (0, 2):
        s = c.page(4, d['title'], 'numbered_editorial_rows')
        y = c.header(s, d['tag'], d['title'], d['subtitle'] if start == 0 else '')
        for n, block in enumerate(d['defenses'][start:start+2]):
            rail_h = text_height(block['title'], 3.2, 18, True) + .65
            body_h = sum(mixed_height(i['full_text'], 8.3, 18, 1.08) + .08
                         for i in block['items'])
            if y + max(rail_h, body_h) > BOTTOM:
                s = c.page(4, d['title'], 'numbered_editorial_rows')
                y = c.header(s, d['tag'], d['title'])
            tx(s, f'{start+n+1:02}', M, y, 3.2, 36, AMBER)
            tx(s, block['title'], M, y + .70, 3.2, 18, bold=True)
            yy = y
            for item in block['items']:
                yy += tx(s, item['full_text'], 4.38, yy, 8.3, 18, MUTED,
                         label_bold=True, leading=1.08) + .08
            y += max(rail_h, body_h) + .22
            if n == 0:
                rule(s, y)
                y += .24

    for idx, ep in enumerate(data['episodes'].values()):
        for side in ('enemy', 'friendly'):
            block = ep[side]
            pending = list(block['bullets'])
            first = True
            while pending:
                s = c.page(ep['source_slide_num'], ep['title'], f'episode_{side}')
                y = c.header(s, ep['tag'], ep['title'], ep['subtitle'] if first else '')
                tx(s, f'2.{idx+1}', M, y - .02, 2.15, 65, AMBER)
                tx(s, block['title'], M, y + 1.15, 2.25, 14, AMBER, True)
                right_y, placed = y, 0
                heights = [mixed_height(it['full_text'], 9.53, 18, 1.10) + .10
                           for it in pending]
                capacity = BOTTOM - y
                page_count = max(1, math.ceil(sum(heights) / capacity))
                target = sum(heights) / page_count if page_count > 1 else capacity
                while pending:
                    item = pending[0]['full_text']
                    needed = mixed_height(item, 9.53, 18, 1.10)
                    if right_y + needed > BOTTOM:
                        break
                    if page_count > 1 and placed >= 2 and right_y - y >= target - .25:
                        break
                    right_y += tx(s, item, 3.15, right_y, 9.53, 18, MUTED,
                                  label_bold=True, leading=1.10) + .10
                    pending.pop(0)
                    placed += 1
                if not placed:
                    raise ValueError('A source paragraph is too long for one page')
                first = False

    d = data['source_8']
    for start in (0, 2):
        s = c.page(8, d['title'], 'matrix_editorial')
        y = c.header(s, d['tag'], d['title'], d['subtitle'] if start == 0 else '')
        widths = [2.35, 2.4, 3.0, 3.53]
        xs = [M, 3.25, 5.9, 9.15]
        header_h = max(text_height(t, w, 12, True) for t, w in zip(d['matrix_headers'], widths))
        for text, x, width in zip(d['matrix_headers'], xs, widths):
            tx(s, text, x, y, width, 12, AMBER, True)
        y += header_h + .15
        rule(s, y)
        y += .22
        for row in d['matrix_rows'][start:start+2]:
            row_h = max(text_height(t, w, 16.5, col == 0)
                        for col, (t, w) in enumerate(zip(row, widths)))
            if y + row_h > BOTTOM:
                raise ValueError('Matrix needs another page')
            for col, (text, x, width) in enumerate(zip(row, xs, widths)):
                tx(s, text, x, y, width, 16.5, WHITE if col == 0 else MUTED, col == 0)
            y += row_h + .25
            rule(s, y)
            y += .25

    s = c.page(8, d['rules_title'], 'closing_statement_rows')
    y = c.header(s, d['tag'], d['rules_title'])
    for i, item in enumerate(d['rules']):
        tx(s, f'0{i+1}', M, y, 1.15, 36, AMBER)
        rh = tx(s, item, 2.1, y + .04, 10.58, 21, leading=1.18)
        y += max(.75, rh) + .28
        if i < 2:
            rule(s, y)
            y += .25
    if y > 6.95:
        raise ValueError('Closing statements overflow')

    c.finish(Presentation(source_path))
    c.prs.save(str(output))
    mapping = {'source_file': Path(source_path).name, 'output_file': output.name,
               'total_slides': len(c.mapping), 'slides': c.mapping}
    map_path = output.parent / 'slide_map.json'
    map_path.write_text(json.dumps(mapping, ensure_ascii=False, indent=2), encoding='utf-8')
    return {'output_path': str(output), 'slide_count': len(c.mapping),
            'slide_map_path': str(map_path), 'min_body_font': '18pt; matrix 16.5pt'}
